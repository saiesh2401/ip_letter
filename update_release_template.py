#!/usr/bin/env python3
"""
Update money_release_template.docx to use separate placeholders:
- {{RELEASE_ORDER}} -> {{COURT_ORDER}} and {{BENEFICIARY}}
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

def update_template():
    template_path = "bank_letters/money_release_template.docx"
    
    try:
        doc = Document(template_path)
        
        # Find and replace the placeholder
        for paragraph in doc.paragraphs:
            if '{{RELEASE_ORDER}}' in paragraph.text:
                # Replace with two separate placeholders
                paragraph.text = paragraph.text.replace(
                    '{{RELEASE_ORDER}}',
                    '{{COURT_ORDER}}\n\n{{BENEFICIARY}}'
                )
                print(f"✅ Updated paragraph: {paragraph.text[:100]}...")
            
            # Also check runs
            for run in paragraph.runs:
                if '{{RELEASE_ORDER}}' in run.text:
                    run.text = run.text.replace(
                        '{{RELEASE_ORDER}}',
                        '{{COURT_ORDER}}\n\n{{BENEFICIARY}}'
                    )
                    print(f"✅ Updated run: {run.text[:100]}...")
        
        # Save the updated template
        doc.save(template_path)
        print(f"\n✅ Successfully updated {template_path}")
        print("   - Replaced {{RELEASE_ORDER}} with {{COURT_ORDER}} and {{BENEFICIARY}}")
        
    except Exception as e:
        print(f"❌ Error updating template: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    update_template()
